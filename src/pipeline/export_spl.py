"""Single Point Lesson deck export (Phase 6f): the VRSI SPL PowerPoint format,
built with SOP Factory's vendored spl_lib.SPLDeck (vendor/sop_factory_2/)
from the same step text, phases and screenshots as every other format.

SOP Factory measures every text box in PowerPoint and refuses to save a deck
with cut-off text. SOPForge has to build a deck on every export, PowerPoint
or not, with no one to ask, so this module:

- plans slides with a font-metric estimate: each slide takes as many of a
  phase's steps (up to MAX_STEPS_PER_SLIDE) as fit at BODY_FLOOR_PT or above,
  with the SOP's own step numbers continuing across slides;
- when PowerPoint is installed, lets spl_lib.fit_text() measure and shrink in
  PowerPoint exactly as SOP Factory does; otherwise shrinks with the same
  estimate (Calibri metrics when the font is on the machine, else the wider
  bundled DejaVu Sans, which errs toward smaller text, never toward overflow);
- never drops or truncates step text: anything still too long at the floor is
  returned as a problem for the review report, and the deck still ships.

No configuration: the PowerPoint path is used when PowerPoint is there."""

import logging
import math
import os
import sys
from pathlib import Path

from PIL import ImageFont
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.util import Pt

from pipeline.assembler import resolve_phases, step_heading
from pipeline.docx_assembler import sop_factory_2_dir
from pipeline.resource_path import resource_path
from pipeline.review_notes import step_review_notes

logger = logging.getLogger(__name__)

TEMPLATE_NAME = "Single Point Lesson Template V4.potm"
MAX_STEPS_PER_SLIDE = 5
BODY_FLOOR_PT = 12
# Template layout placeholders (Assembly Instruction x1/x2 Images).
TASK_IDX, STEPS_IDX, NOTES_IDX = 2, 13, 14
# PowerPoint text-frame defaults (bodyPr insets) and the template's body
# paragraph rules (master bodyStyle: 90% line spacing, 7.5pt before).
INSET_X_PT, INSET_Y_PT = 7.2, 3.6
LINE_HEIGHT_EM = 1.22 * 0.9
PARA_SPACE_PT = 7.5
HEIGHT_SAFETY = 1.05
_EMU_PER_PT = 12700
_TEMPLATE_FURNITURE = {PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER}


def _import_spl():
    vendor = str(sop_factory_2_dir())
    if vendor not in sys.path:
        sys.path.insert(0, vendor)
    import spl_lib

    return spl_lib


def _font_files():
    """(regular, bold) font files to measure with: Calibri, the template's
    theme font, when installed; else the bundled DejaVu Sans."""
    windir = os.environ.get("WINDIR")
    if windir:
        fonts = Path(windir) / "Fonts"
        regular, bold = fonts / "calibri.ttf", fonts / "calibrib.ttf"
        if regular.exists() and bold.exists():
            return regular, bold
    dejavu = resource_path("assets", "fonts", "dejavu-sans")
    return dejavu / "DejaVuSans.ttf", dejavu / "DejaVuSans-Bold.ttf"


class TextEstimator:
    """Greedy word-wrap height estimate for a PowerPoint text box."""

    def __init__(self):
        self._files = _font_files()
        self._fonts = {}

    def _font(self, bold):
        if bold not in self._fonts:
            # Measured at 100 units and scaled: widths are linear in size.
            self._fonts[bold] = ImageFont.truetype(str(self._files[bold]), 100)
        return self._fonts[bold]

    def _lines(self, text, size, width, bold):
        """Lines the text wraps to at this size and width: greedy word wrap,
        with a word wider than the box broken across as many lines as it
        needs (PowerPoint breaks it mid-word)."""
        if "\n" in text or "\v" in text:
            return sum(self._lines(line, size, width, bold) for line in text.splitlines())
        font = self._font(bold)
        scale = size / 100
        space = font.getlength(" ") * scale
        lines, used = 1, 0.0
        for word in text.split():
            w = font.getlength(word) * scale
            if used and used + space + w <= width:
                used += space + w
                continue
            if used:
                lines += 1
            extra = max(0, math.ceil(w / width) - 1)
            lines += extra
            used = w - extra * width
        return lines

    def fits(self, box_pt, paragraphs, size, bold=False, indent_pt=0.0, para_space=0.0):
        width = box_pt[0] - 2 * INSET_X_PT - indent_pt
        height = box_pt[1] - 2 * INSET_Y_PT
        if width <= 0 or height <= 0:
            return False
        total = sum(self._lines(p, size, width, bold) * size * LINE_HEIGHT_EM for p in paragraphs)
        total += para_space * (len(paragraphs) - 1)  # none before a frame's first line
        return total * HEIGHT_SAFETY <= height

    def fit_size(self, box_pt, paragraphs, base, floor, **kw):
        """Largest whole point size from base down to floor that fits, or None."""
        size = base
        while size >= floor:
            if self.fits(box_pt, paragraphs, size, **kw):
                return size
            size -= 1
        return None


def _box_pt(placeholder):
    return placeholder.width / _EMU_PER_PT, placeholder.height / _EMU_PER_PT


def _lvl1(placeholder):
    found = placeholder._element.xpath("./p:txBody/a:lstStyle/a:lvl1pPr")
    return found[0] if found else None


def _text_rules(slide_placeholder, master):
    """(base size, bold, indent pt, paragraph space pt) for a slide
    placeholder, from its layout placeholder, else the master's styles."""
    is_title = slide_placeholder.placeholder_format.idx == 0
    base_ph = getattr(slide_placeholder, "_base_placeholder", None)
    lvl = _lvl1(base_ph) if base_ph is not None else None
    size = bold = indent = None
    if lvl is not None:
        defrpr = lvl.find("{http://schemas.openxmlformats.org/drawingml/2006/main}defRPr")
        if defrpr is not None and defrpr.get("sz"):
            size = int(defrpr.get("sz")) / 100
        if defrpr is not None:
            bold = defrpr.get("b") == "1"
        if lvl.get("marL"):
            indent = int(lvl.get("marL")) / _EMU_PER_PT
    if size is None:
        style = "p:titleStyle" if is_title else "p:bodyStyle"
        sz = master.element.xpath(f".//{style}/a:lvl1pPr/a:defRPr/@sz")
        size = int(sz[0]) / 100 if sz else 18
    return size, bool(bold), indent or 0.0, 0.0 if is_title else PARA_SPACE_PT


def estimate_fit(pptx_path, estimator=None):
    """Shrinks overflowing placeholder text in a saved deck by the estimate,
    1pt at a time down to the floor (12pt, or 75% of a smaller base), saves
    it, and returns the boxes that still overflow -- the no-PowerPoint
    counterpart of spl_lib.fit_text."""
    estimator = estimator or TextEstimator()
    prs = Presentation(str(pptx_path))
    problems = []
    for n, slide in enumerate(prs.slides, start=1):
        for ph in slide.placeholders:
            if not ph.has_text_frame or not ph.text_frame.text.strip():
                continue
            if ph.placeholder_format.type in _TEMPLATE_FURNITURE:
                continue  # date / footer / slide number: short, fixed template text
            base, bold, indent, para_space = _text_rules(ph, prs.slide_master)
            bold = bold or any(r.font.bold for p in ph.text_frame.paragraphs for r in p.runs)
            run_sizes = [
                r.font.size.pt for p in ph.text_frame.paragraphs for r in p.runs if r.font.size
            ]
            base = min(run_sizes) if run_sizes else base
            floor = BODY_FLOOR_PT if base >= BODY_FLOOR_PT else base * 0.75
            paragraphs = [p.text for p in ph.text_frame.paragraphs]
            kw = {"bold": bold, "indent_pt": indent, "para_space": para_space}
            size = estimator.fit_size(_box_pt(ph), paragraphs, base, floor, **kw)
            if size is None:
                problems.append(f"slide {n} '{ph.name}': {ph.text_frame.text[:60]!r}")
                size = floor
            if size != base or run_sizes:
                for p in ph.text_frame.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(size)
    prs.save(str(pptx_path))
    return problems


def powerpoint_available():
    """True when PowerPoint's COM class is registered (Windows only); checked
    in the registry, so nothing is launched just to find out."""
    if sys.platform != "win32":
        return False
    try:
        import winreg

        winreg.CloseKey(winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, "PowerPoint.Application"))
        return True
    except OSError:
        return False


def _fit_in_powerpoint(spl_lib, pptx_path):
    import pythoncom  # export jobs run on worker threads: COM must be initialized per thread

    pythoncom.CoInitialize()
    try:
        return spl_lib.fit_text(Path(pptx_path))
    finally:
        pythoncom.CoUninitialize()


def _step_text(n, step, result):
    # Preserve the compact capture layout; photo titles are part of the SOP
    # and must also be included in slide planning's fit calculation.
    if result.get("title"):
        return f"{step_heading(n, step, result)}\v{result['text']}"
    return result["text"]


def plan_slides(manifest, step_results, phases=None, estimator=None, steps_box_pt=(690, 107)):
    """[(phase name, [(step number, step, result), ...]), ...]: each phase's
    steps packed onto as few slides as fit at the floor, in order, every
    step exactly once."""
    estimator = estimator or TextEstimator()
    numbered = {
        step.id: (n, step, result)
        for n, (step, result) in enumerate(zip(manifest.steps, step_results, strict=True), start=1)
    }
    kw = {"indent_pt": 342900 / _EMU_PER_PT, "para_space": PARA_SPACE_PT}
    slides = []
    for phase in resolve_phases(manifest, phases):
        current = []
        for step_id in phase.step_ids:
            candidate = current + [numbered[step_id]]
            texts = [_step_text(n, s, r) for n, s, r in candidate]
            if len(candidate) <= MAX_STEPS_PER_SLIDE and estimator.fits(
                steps_box_pt,
                texts,
                BODY_FLOOR_PT,
                bold=any(r.get("title") for _, _, r in candidate),
                **kw,
            ):
                current = candidate
            else:
                if current:
                    slides.append((phase.name, current))
                current = [numbered[step_id]]
        if current:
            slides.append((phase.name, current))
    return slides


def _notes_line(entries):
    notes = [note for _n, step, _r in entries for note in step_review_notes(step)]
    if len(notes) == 1:
        return notes[0]
    if notes:
        return f"{len(notes)} items need review: see the SOP."
    return ""


def _task_label(name, continued, estimator, box_pt):
    """The slide's bold task label: the phase name, "(cont.)" on later
    slides, shortened with an ellipsis only if it can't fit at the floor.
    It names a section, not a step, so shortening it loses no step text."""

    def fits(label):
        return estimator.fit_size(box_pt, [label], 18, BODY_FLOOR_PT, bold=True) is not None

    label = f"{name} (cont.)" if continued else name
    words = label.split()
    while not fits(label) and len(words) > 1:
        words.pop()
        label = " ".join(words) + "…"
    return label


def build_spl_deck(
    manifest,
    step_results,
    annotated_paths,
    output_path,
    author="SOPForge",
    phases=None,
    omitted_screenshots=(),
):
    """Builds the SPL .pptx at output_path. Returns (output_path, problems):
    problems lists text boxes that still overflow at the floor size (the deck
    is written either way)."""
    spl_lib = _import_spl()
    estimator = TextEstimator()
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    deck = spl_lib.SPLDeck(
        str(output_path),
        author=author,
        active_dir=str(output_path.parent),
        template=sop_factory_2_dir() / TEMPLATE_NAME,
    )
    shot_by_id = {
        step.id: Path(shot)
        for step, shot in zip(manifest.steps, annotated_paths, strict=True)
        if shot is not None and Path(shot).exists() and step.id not in omitted_screenshots
    }
    title = manifest.session.title or manifest.session.id
    first_shot = next(iter(shot_by_id.values()), None)
    deck.title_slide(title, image=str(first_shot) if first_shot else None)

    layout = deck.layouts["Assembly Instruction x1 Image"]
    task_box = _box_pt(layout.placeholders.get(idx=TASK_IDX))
    steps_box = _box_pt(layout.placeholders.get(idx=STEPS_IDX))
    seen_phases = set()
    for phase_name, entries in plan_slides(manifest, step_results, phases, estimator, steps_box):
        images = [str(shot_by_id[s.id]) for _n, s, _r in entries if s.id in shot_by_id][:2]
        slide = deck.step(
            _task_label(phase_name, phase_name in seen_phases, estimator, task_box),
            [_step_text(n, s, r) for n, s, r in entries],
            notes=_notes_line(entries),
            images=images,
            start_at=entries[0][0],
        )
        body = slide.placeholders[STEPS_IDX]
        for paragraph, (n, step, result) in zip(body.text_frame.paragraphs, entries, strict=True):
            if result.get("title"):
                paragraph.clear()
                run = paragraph.add_run()
                run.text = step_heading(n, step, result)
                run.font.bold = True
                paragraph.add_line_break()
                paragraph.add_run().text = result["text"]
        seen_phases.add(phase_name)

    # What spl_lib.save() does before fitting, minus the full local path: the
    # template's footer shows the file name, which is all a recipient needs.
    for shape in deck.prs.slide_master.shapes:
        if shape.name == "FilePath":
            shape.text_frame.paragraphs[0].runs[0].text = output_path.name
    deck.prs.save(str(output_path))

    if powerpoint_available():
        try:
            return output_path, _fit_in_powerpoint(spl_lib, output_path)
        except Exception as exc:  # noqa: BLE001 - fall back to the estimate, never fail the export
            logger.warning("PowerPoint text fit failed (%s); using the estimate", exc)
    return output_path, estimate_fit(output_path, estimator)
