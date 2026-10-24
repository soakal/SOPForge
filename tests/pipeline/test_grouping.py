"""Phase 6c (phases/06-factory-parity.md AC3): phase grouping never drops,
duplicates or reorders a step, whichever path chose the phases (LLM accepted,
LLM rejected -> window fallback, fallback only), and only exact-duplicate
screenshots are omitted, each one listed in the sidecar report."""

import copy
import json
import time
from pathlib import Path

import docx
from fastapi.testclient import TestClient
from hypothesis import given
from hypothesis import strategies as st
from PIL import Image, ImageDraw

from pipeline.assembler import toc_lines
from pipeline.docx_assembler import assemble_docx
from pipeline.export_html import render_single_file_html
from pipeline.grouping import (
    DEFAULT_PHASE_NAME,
    Phase,
    default_phases,
    duplicate_screenshot_steps,
    group_by_window,
    load_keep_screenshots,
    load_phases,
    parse_proposal,
    phases_to_json,
    propose_phases,
    validate_phases,
)
from pipeline.manifest import Manifest
from pipeline.render import render_html, render_markdown, render_steps_template_mode
from pipeline.server import create_app

FIXTURES = Path(__file__).resolve().parent.parent.parent / "fixtures"
SIGN_IN = ("SmartDeploy.exe", "Sign in - SmartDeploy Console")
EDITOR = ("SmartDeploy.exe", "Answer File Editor - SmartDeploy Console")
BROWSER = ("chrome.exe", "Downloads - Google Chrome")


def _manifest_dict(windows):
    """sample-manifest's steps cycled to len(windows), one window each."""
    base = json.loads((FIXTURES / "sample-manifest.json").read_text(encoding="utf-8"))
    template = [s for s in base["steps"] if s["action"] == "click" and s["element"]["name"]][0]
    steps = []
    for n, (process, title) in enumerate(windows, start=1):
        step = copy.deepcopy(template)
        step["id"] = f"step-{n:03d}"
        step["screenshot"] = f"{n:03d}.png"
        step["window"]["process"], step["window"]["title"] = process, title
        step["element"]["name"] = f"Button {n}"
        step["element"]["automation_id"] = f"btn{n}"
        steps.append(step)
    base["steps"] = steps
    return base


def _manifest(windows):
    return Manifest.model_validate(_manifest_dict(windows))


SIX = [SIGN_IN, SIGN_IN, EDITOR, BROWSER, EDITOR, EDITOR]


class _ScriptedLLM:
    def __init__(self, reply):
        self.reply = reply
        self.calls = 0

    def chat(self, messages, **kwargs):
        self.calls += 1
        if isinstance(self.reply, Exception):
            raise self.reply
        return self.reply


# ---- validation ---------------------------------------------------------------


def test_validate_accepts_exact_partition_and_rejects_everything_else():
    m = _manifest(SIX)
    ids = [s.id for s in m.steps]
    assert validate_phases([Phase("A", tuple(ids[:2])), Phase("B", tuple(ids[2:]))], m)
    bad = [
        [],
        [Phase("A", tuple(ids[:2]))],  # drops steps
        [Phase("A", tuple(ids[:3])), Phase("B", tuple(ids[2:]))],  # duplicate
        [Phase("A", tuple(ids[2:])), Phase("B", tuple(ids[:2]))],  # reordered
        [Phase("A", (ids[0], ids[2])), Phase("B", (ids[1], *ids[3:]))],  # non-contiguous
        [Phase(" ", tuple(ids))],  # blank name
        [Phase("A", tuple(ids[:2])), Phase("B", ()), Phase("C", tuple(ids[2:]))],  # empty
        [Phase("x" * 61, tuple(ids))],  # name too long
    ]
    for phases in bad:
        assert not validate_phases(phases, m), phases


@given(st.lists(st.integers(min_value=0, max_value=5), min_size=0, max_size=8))
def test_any_permuted_or_dropped_grouping_is_rejected(picks):
    """Property: phases built from any id sequence other than the exact
    manifest order (dropped, repeated or shuffled ids) never validate."""
    m = _manifest(SIX)
    ids = [s.id for s in m.steps]
    chosen = [ids[i] for i in picks]
    phases = [Phase("A", tuple(chosen))] if chosen else []
    assert validate_phases(phases, m) == (chosen == ids)


# ---- LLM proposal ---------------------------------------------------------------


def test_parse_proposal_builds_contiguous_phases():
    m = _manifest(SIX)
    phases = parse_proposal(
        'Sure: [{"name": "Sign in", "start": 1}, {"name": "Edit  the answer file", "start": 3}]',
        m,
    )
    assert [p.name for p in phases] == ["Sign in", "Edit the answer file"]
    assert phases[0].step_ids == ("step-001", "step-002")
    assert phases[1].step_ids == tuple(f"step-{n:03d}" for n in range(3, 7))


def test_parse_proposal_rejects_bad_replies():
    m = _manifest(SIX)
    for reply in (
        "no json here",
        "[not json]",
        "[]",
        '[{"name": "A", "start": 2}]',  # doesn't start at step 1
        '[{"name": "A", "start": 1}, {"name": "B", "start": 1}]',  # not increasing
        '[{"name": "A", "start": 1}, {"name": "B", "start": 7}]',  # past the end
        '[{"name": "A", "start": true}]',
        '[{"name": 5, "start": 1}]',
        '[{"name": "' + "x" * 61 + '", "start": 1}]',
    ):
        assert parse_proposal(reply, m) is None, reply


def test_propose_uses_llm_when_valid_and_falls_back_otherwise():
    m = _manifest(SIX)
    good = _ScriptedLLM('[{"name": "Sign in", "start": 1}, {"name": "Configure", "start": 3}]')
    phases, source = propose_phases(m, [], good)
    assert source == "llm" and [p.name for p in phases] == ["Sign in", "Configure"]
    for llm in (_ScriptedLLM("garbage"), _ScriptedLLM(RuntimeError("down")), None):
        phases, source = propose_phases(m, [], llm)
        assert source == "window"
        assert validate_phases(phases, m)


def test_short_captures_stay_ungrouped_without_an_llm_call():
    m = _manifest(SIX[:3])
    llm = _ScriptedLLM('[{"name": "A", "start": 1}]')
    assert propose_phases(m, [], llm) == (default_phases(m), "default")
    assert llm.calls == 0


# ---- window fallback --------------------------------------------------------------


def test_group_by_window_folds_one_step_detours():
    phases = group_by_window(_manifest(SIX))
    assert [p.name for p in phases] == [
        "Sign in - SmartDeploy Console",
        "Answer File Editor - SmartDeploy Console",
    ]
    assert phases[1].step_ids == ("step-003", "step-004", "step-005", "step-006")


def test_group_by_window_single_window_is_default_and_strips_browser_suffix():
    assert group_by_window(_manifest([EDITOR] * 5)) == default_phases(_manifest([EDITOR] * 5))
    phases = group_by_window(_manifest([BROWSER, BROWSER, EDITOR, EDITOR]))
    assert phases[0].name == "Downloads"


def test_load_phases_rejects_stale_file(tmp_path):
    m = _manifest(SIX)
    path = tmp_path / "phases.json"
    saved = [Phase("A", ("step-001", "step-002")), Phase("B", tuple(s.id for s in m.steps[2:]))]
    path.write_text(json.dumps(phases_to_json(saved, "llm")), encoding="utf-8")
    assert load_phases(path, m) == saved
    shorter = _manifest(SIX[:5])  # a step was hidden since the file was written
    assert load_phases(path, shorter) == group_by_window(shorter)
    path.write_text("{broken", encoding="utf-8")
    assert load_phases(path, m) == group_by_window(m)


# ---- duplicate screenshots ----------------------------------------------------------


def _shot(path, marker_x):
    img = Image.new("RGB", (400, 200), (255, 255, 255))
    ImageDraw.Draw(img).ellipse((marker_x, 80, marker_x + 30, 110), outline=(255, 0, 0), width=4)
    img.save(path)
    return path


def test_only_exact_repeats_of_the_same_control_are_omitted(tmp_path):
    data = _manifest_dict([EDITOR] * 5)
    for step in data["steps"][:3]:  # steps 1-3 click the same control
        step["element"]["name"], step["element"]["automation_id"] = "Save", "save"
    m = Manifest.model_validate(data)
    paths = [
        _shot(tmp_path / "1.png", 100),
        _shot(tmp_path / "2.png", 100),  # identical repeat of step 1 -> omitted
        _shot(tmp_path / "3.png", 250),  # same control, different picture -> kept
        _shot(tmp_path / "4.png", 250),  # same picture, different control -> kept
        _shot(tmp_path / "5.png", 250),
    ]
    one = [Phase("All", tuple(s.id for s in m.steps))]
    assert duplicate_screenshot_steps(m, paths, one) == ["step-002"]
    split = [Phase("A", ("step-001",)), Phase("B", tuple(s.id for s in m.steps[1:]))]
    assert duplicate_screenshot_steps(m, paths, split) == []  # a phase's first step keeps it


# ---- rendering --------------------------------------------------------------------


def _rendered(tmp_path, windows=SIX):
    m = _manifest(windows)
    shots = tmp_path / "shots"
    shots.mkdir()
    for step in m.steps:
        Image.new("RGB", (1920, 1080), (255, 255, 255)).save(shots / step.screenshot)
    results, paths = render_steps_template_mode(m, shots, tmp_path / "annotated")
    phases = [Phase("Sign in", ("step-001", "step-002")), Phase("Configure", m.step_ids()[2:])]
    return m, results, paths, tmp_path / "annotated", phases


def test_toc_numbers_each_phase():
    m = _manifest(SIX)
    phases = [Phase("Sign in", ("step-001", "step-002")), Phase("Configure", m.step_ids()[2:])]
    lines = toc_lines(m, "overview", phases)
    assert lines[:3] == ["1.  Overview", "2.  Sign in", "      Step 1 — Click 'Button 1'"]
    assert "4.  Revision History" == lines[-1] and "3.  Configure" in lines
    assert toc_lines(m, None)[0] == f"1.  {DEFAULT_PHASE_NAME}"


def test_grouped_docx_has_phase_sections_and_every_step(tmp_path):
    m, results, _, annotated, phases = _rendered(tmp_path)
    out, _ = assemble_docx(
        m,
        results,
        annotated,
        tmp_path / "out.docx",
        phases=phases,
        omitted_screenshots={"step-004"},
    )
    document = docx.Document(out)
    texts = [p.text for p in document.paragraphs]
    assert "Sign in" in texts and "Configure" in texts and "Procedure" not in texts
    assert texts.index("Sign in") < texts.index("Configure")
    for n, step in enumerate(m.steps, start=1):
        assert any(t.startswith(f"Step {n} —") for t in texts), step.id
    pictures = sum(1 for p in document.paragraphs if p._p.xpath(".//w:drawing"))
    assert pictures == len(m.steps) - 1


def test_grouped_markdown_and_html(tmp_path):
    m, results, paths, _, phases = _rendered(tmp_path)
    omitted = {"step-004"}
    md = render_markdown(m, results, paths, phases=phases, omitted_screenshots=omitted)
    assert "## Sign in" in md and "## Configure" in md and "### Step 1 —" in md
    assert md.count("![") == len(m.steps) - 1
    page = render_html(m, results, paths, phases=phases, omitted_screenshots=omitted)
    assert "<h2>Sign in</h2>" in page
    for step in m.steps:  # review-page anchors survive grouping
        assert f'<h3 id="{step.id}">' in page
    assert page.count("<img ") == len(m.steps) - 1
    single = render_single_file_html(m, results, paths, phases=phases, omitted_screenshots=omitted)
    assert "<h2>Configure</h2>" in single and single.count("<img ") == len(m.steps) - 1


def test_ungrouped_markdown_is_unchanged(tmp_path):
    m, results, paths, _, _ = _rendered(tmp_path)
    assert render_markdown(m, results, paths) == render_markdown(
        m, results, paths, phases=default_phases(m)
    )


# ---- end to end through the server --------------------------------------------------


def _phase_aware_factory(reply):
    class Client:
        def chat(self, messages, **kwargs):
            if "Split them into" in messages[-1]["content"]:
                return reply
            return "stub reply that never matches any manifest, forcing template fallback"

    return lambda: Client()


def _run_session(tmp_path, reply):
    factory = _phase_aware_factory(reply)
    client = TestClient(
        create_app(
            sessions_root=tmp_path / "sessions",
            llm_client_factory=factory,
            narrative_llm_client_factory=factory,
        )
    )
    data = _manifest_dict(SIX)
    shots = tmp_path / "upload"
    shots.mkdir()
    files = []
    for step in data["steps"]:
        p = shots / step["screenshot"]
        Image.new("RGB", (1920, 1080), (255, 255, 255)).save(p)
        files.append(("files", (step["screenshot"], p.open("rb"), "image/png")))
    resp = client.post("/sessions", data={"manifest_json": json.dumps(data)}, files=files)
    assert resp.status_code == 200
    session_id = resp.json()["session_id"]
    _wait(client, session_id)
    return client, session_id


def _wait(client, session_id):
    deadline = time.monotonic() + 15
    while time.monotonic() < deadline:
        status = client.get(f"/sessions/{session_id}/status").json()["status"]
        if status in ("done", "error"):
            assert status == "done"
            return
        time.sleep(0.05)
    raise AssertionError("session never finished")


def test_server_applies_llm_phases_and_keeps_them_on_reexport(tmp_path):
    reply = '[{"name": "Sign in", "start": 1}, {"name": "Configure", "start": 3}]'
    client, session_id = _run_session(tmp_path, reply)
    report = client.get(f"/sessions/{session_id}/report").json()
    assert [p["name"] for p in report["phases"]] == ["Sign in", "Configure"]
    flat = [i for p in report["phases"] for i in p["step_ids"]]
    assert flat == [f"step-{n:03d}" for n in range(1, 7)]
    assert "## Sign in" in client.get(f"/sessions/{session_id}/doc.md").text

    # A step edit re-exports from disk with no LLM call; phases must survive.
    resp = client.post(
        f"/ui/sessions/{session_id}/steps/step-003",
        data={"text": "Click Button 3 in the Answer File Editor window."},
        follow_redirects=False,
    )
    assert resp.status_code in (200, 303)
    _wait(client, session_id)
    md = client.get(f"/sessions/{session_id}/doc.md").text
    assert "## Sign in" in md and "## Configure" in md


def test_server_falls_back_to_window_grouping_on_a_bad_proposal(tmp_path):
    client, session_id = _run_session(tmp_path, "I can't do that.")
    report = client.get(f"/sessions/{session_id}/report").json()
    assert [p["name"] for p in report["phases"]] == [
        "Sign in - SmartDeploy Console",
        "Answer File Editor - SmartDeploy Console",
    ]


def test_restored_screenshots_round_trip_through_phases_json(tmp_path):
    m = _manifest(SIX)
    path = tmp_path / "phases.json"
    path.write_text(
        json.dumps(phases_to_json(default_phases(m), "user", keep_screenshots={"step-004"})),
        encoding="utf-8",
    )
    assert load_keep_screenshots(path) == {"step-004"}
    assert load_phases(path, m) == default_phases(m)
    path.write_text("{broken", encoding="utf-8")
    assert load_keep_screenshots(path) == frozenset()
    assert "keep_screenshots" not in phases_to_json(default_phases(m), "llm")
