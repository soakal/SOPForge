"""Phase 6f (phases/06-factory-parity.md AC6): the SPL deck always builds,
with or without PowerPoint, carries every step's full text in order with the
SOP's own numbering, and reports (never hides) text that can't fit."""

import copy
import json
from pathlib import Path

from PIL import Image
from pptx import Presentation

import pipeline.export_spl as export_spl
from pipeline.export_spl import (
    BODY_FLOOR_PT,
    MAX_STEPS_PER_SLIDE,
    TextEstimator,
    build_spl_deck,
    plan_slides,
)
from pipeline.grouping import Phase
from pipeline.manifest import Manifest
from pipeline.render import render_steps_template_mode

FIXTURES = Path(__file__).resolve().parent.parent.parent / "fixtures"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _manifest(n):
    base = json.loads((FIXTURES / "sample-manifest.json").read_text(encoding="utf-8"))
    template = base["steps"][0]
    steps = []
    for i in range(1, n + 1):
        step = copy.deepcopy(template)
        step["id"], step["screenshot"] = f"step-{i:03d}", f"{i:03d}.png"
        step["element"]["name"] = f"Button {i}"
        steps.append(step)
    base["steps"] = steps
    base["session"]["title"] = "Sample procedure"
    return Manifest.model_validate(base)


def _rendered(tmp_path, n=9):
    m = _manifest(n)
    shots = tmp_path / "shots"
    shots.mkdir()
    for step in m.steps:
        Image.new("RGB", (1920, 1080), (255, 255, 255)).save(shots / step.screenshot)
    results, paths = render_steps_template_mode(m, shots, tmp_path / "annotated")
    return m, results, paths


def _slide_texts(path):
    prs = Presentation(str(path))
    return [
        [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text]
        for slide in prs.slides
    ]


def test_estimator_fits_short_text_and_rejects_a_wall():
    est = TextEstimator()
    assert est.fits((690, 107), ["Click Save."], 18)
    assert not est.fits((690, 107), ["word " * 400], BODY_FLOOR_PT)
    assert est.fit_size((690, 107), ["word " * 400], 18, BODY_FLOOR_PT) is None


def test_estimator_breaks_an_overlong_word_across_lines():
    est = TextEstimator()
    assert est._lines("x" * 400, 18, 100, False) > 3


def test_plan_keeps_every_step_once_in_order_within_phases():
    m = _manifest(9)
    results = [{"step_id": s.id, "text": f"Click Button {i}."} for i, s in enumerate(m.steps, 1)]
    phases = [Phase("A", m.step_ids()[:2]), Phase("B", m.step_ids()[2:])]
    slides = plan_slides(m, results, phases)
    numbers = [n for _name, entries in slides for n, _s, _r in entries]
    assert numbers == list(range(1, 10))
    assert all(len(entries) <= MAX_STEPS_PER_SLIDE for _name, entries in slides)
    assert [name for name, _ in slides][:1] == ["A"] and {n for n, _, _ in slides[0][1]} == {1, 2}


def test_plan_gives_an_overlong_step_its_own_slide():
    m = _manifest(3)
    results = [
        {"step_id": "step-001", "text": "Short."},
        {"step_id": "step-002", "text": "word " * 300},
        {"step_id": "step-003", "text": "Short."},
    ]
    slides = plan_slides(m, results)
    assert [[n for n, _s, _r in e] for _name, e in slides] == [[1], [2], [3]]


def test_deck_has_every_step_with_continuing_numbers(tmp_path):
    m, results, paths = _rendered(tmp_path)
    phases = [Phase("Sign in", m.step_ids()[:2]), Phase("Configure", m.step_ids()[2:])]
    out, problems = build_spl_deck(m, results, paths, tmp_path / "deck.pptx", phases=phases)
    assert problems == []
    texts = _slide_texts(out)
    assert texts[0][0] == "Sample procedure"  # title slide
    body = "\n".join(t for slide in texts[1:] for t in slide)
    for result in results:
        assert result["text"] in body  # full text, never truncated
    assert any("Configure (cont.)" in t for slide in texts for t in slide)
    prs = Presentation(str(out))
    starts = [
        int(el.get("startAt"))
        for slide in prs.slides
        for el in slide.shapes._spTree.iter(f"{{{A}}}buAutoNum")
        if el.get("startAt")
    ]
    assert starts and starts == sorted(starts) and starts[0] > 1
    pictures = sum(1 for slide in prs.slides for sh in slide.shapes if sh.shape_type == 13)
    assert pictures >= len(prs.slides)  # title + at least one per step slide


def test_text_too_long_is_reported_not_cut(tmp_path):
    m, results, paths = _rendered(tmp_path, n=1)
    results[0]["text"] = "word " * 300
    out, problems = build_spl_deck(m, results, paths, tmp_path / "deck.pptx")
    assert out.exists() and problems
    assert ("word " * 300).strip() in "\n".join(t for s in _slide_texts(out) for t in s)


def test_powerpoint_failure_falls_back(tmp_path, monkeypatch):
    monkeypatch.setattr(export_spl, "powerpoint_available", lambda: True)

    def boom(*_a):
        raise RuntimeError("COM server busy")

    monkeypatch.setattr(export_spl, "_fit_in_powerpoint", boom)
    m, results, paths = _rendered(tmp_path, n=3)
    out, problems = build_spl_deck(m, results, paths, tmp_path / "deck.pptx")
    assert out.exists() and problems == []


def test_footer_shows_file_name_not_local_path(tmp_path):
    m, results, paths = _rendered(tmp_path, n=2)
    out, _ = build_spl_deck(m, results, paths, tmp_path / "deck.pptx")
    master = Presentation(str(out)).slide_master
    file_path_text = [s.text_frame.text for s in master.shapes if s.name == "FilePath"]
    assert file_path_text == ["deck.pptx"]


def test_photo_titles_survive_slide_export_and_fit_planning(tmp_path, monkeypatch):
    monkeypatch.setattr(export_spl, "powerpoint_available", lambda: False)
    m, results, paths = _rendered(tmp_path, n=3)
    results[0]["title"] = "Select <Products> & licenses"
    out, _ = build_spl_deck(m, results, paths, tmp_path / "titled.pptx")
    text = "\n".join(t for slide in _slide_texts(out) for t in slide)
    assert "Step 1 — Select <Products> & licenses" in text
    assert any(
        r.font.bold and r.text == "Step 1 — Select <Products> & licenses"
        for slide in Presentation(out).slides
        for shape in slide.shapes
        if shape.has_text_frame
        for paragraph in shape.text_frame.paragraphs
        for r in paragraph.runs
    )
    assert results[0]["text"] in text

    class Estimator:
        def fits(self, box, texts, *args, **kwargs):
            assert any("Select <Products>" in text for text in texts)
            return True

    plan_slides(m, results, estimator=Estimator())
